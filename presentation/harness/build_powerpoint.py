"""Build the fidelity-preserving PowerPoint with notes and an embedded demo.

First run `node presentation/harness/export.mjs --powerpoint`.
Artwork is rasterized at 2400 x 1350. Notes and the demo link remain editable.
"""
from __future__ import annotations

import json
import re
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import imageio_ffmpeg
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
EXPORTS = ROOT / ".exports"
TITLES = [
    "What it is and who it is for",
    "What goes in and what comes out",
    "Why this approach fits",
    "Why Flower",
    "The embedded gateway demonstration",
    "The next experiment",
]


def build():
    for index in range(1, 7):
        if not (EXPORTS / f"slide-{index:02}.png").is_file():
            raise SystemExit("Run export.mjs --powerpoint before building PowerPoint.")

    frames = EXPORTS / "frames.txt"
    lines = []
    for index in range(11):
        lines.extend([f"file 'demo-{index:02}.png'", f"duration {2 if index == 0 else 8 if index == 10 else 3}"])
    lines.append("file 'demo-10.png'")
    frames.write_text("\n".join(lines) + "\n", encoding="utf-8")
    video = EXPORTS / "gateway-demo.mp4"
    subprocess.run([
        imageio_ffmpeg.get_ffmpeg_exe(), "-y", "-loglevel", "error",
        "-f", "concat", "-safe", "0", "-i", str(frames),
        "-t", "37", "-vf", "fps=30", "-c:v", "libx264", "-preset", "medium",
        "-crf", "18", "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(video),
    ], check=True)

    speech = (ROOT / "speech_for_harness.txt").read_text(encoding="utf-8")
    sections = re.split(r"\n\d+\. [^\n]+\n", speech)[1:]
    assert len(sections) == 6
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(40 / 3), Inches(7.5)
    deck.core_properties.title = "Fusion Investigator — Flower agent harness"
    deck.core_properties.subject = "Six-slide, five-minute talk with embedded synthetic gateway replay"
    deck.core_properties.author = "Fusion Investigator project"
    deck.core_properties.keywords = "Flower, fusion, gateway, synthetic, investigation"

    for index, (title, section) in enumerate(zip(TITLES, sections), start=1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.name = title
        picture = slide.shapes.add_picture(str(EXPORTS / f"slide-{index:02}.png"), 0, 0, deck.slide_width, deck.slide_height)
        picture.name = f"Slide {index} artwork — {title}"
        picture._element.xpath(".//p:cNvPr")[0].set("descr", re.sub(r"\[[^]]*\]", "", section).strip())
        slide.notes_slide.notes_text_frame.text = section.strip()
        if index == 5:
            movie = slide.shapes.add_movie(str(video), 0, 0, deck.slide_width, deck.slide_height,
                                         poster_frame_image=str(EXPORTS / "slide-05.png"), mime_type="video/mp4")
            movie.name = "Gateway replay — click to play — 37 seconds — no LLM"
            label = slide.shapes.add_textbox(Inches(.69), Inches(6.64), Inches(7.4), Inches(.26))
            paragraph = label.text_frame.paragraphs[0]
            paragraph.text = "Click the chart to play  /  37-second gateway replay  /  no model calls"
            paragraph.font.name = "Arial"
            paragraph.font.size = Pt(11)
            paragraph.font.color.rgb = RGBColor.from_string("A0ACB8")
            link = slide.shapes.add_textbox(Inches(10.56), Inches(6.61), Inches(2.2), Inches(.34))
            run = link.text_frame.paragraphs[0].add_run()
            run.text = "Open interactive demo ↗"
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor.from_string("FF986E")
            link.click_action.hyperlink.address = "index.html#investigation"
            slide.notes_slide.notes_text_frame.text += (
                "\n\nPOWERPOINT: Click the chart to play the embedded 37-second MP4. "
                "Narrate the ten gateway results as the video advances. "
                "The video is embedded and works offline. For request-by-request "
                "interaction, use the linked browser deck in the same folder."
            )

    output = ROOT / "fusion-investigator.pptx"
    deck.save(output)
    # Check the actual OOXML package, including embedded media, relationships and notes.
    ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    with zipfile.ZipFile(output) as package:
        assert package.testzip() is None
        names = package.namelist()
        slides = [name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)]
        notes = [name for name in names if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)]
        assert len(slides) == len(notes) == 6
        assert any(name.endswith(".mp4") and name.startswith("ppt/media/") for name in names)
        demo = ET.fromstring(package.read("ppt/slides/slide5.xml"))
        assert demo.find(".//a:videoFile", ns) is not None
        rels = package.read("ppt/slides/_rels/slide5.xml.rels").decode()
        assert "index.html#investigation" in rels
    reopened = Presentation(output)
    assert len(reopened.slides) == 6
    assert "Flower's agent harness" in reopened.slides[0].notes_slide.notes_text_frame.text
    print(json.dumps({"output": str(output), "slides": 6, "speaker_notes": 6,
                      "embedded_demo_seconds": 37, "size_bytes": output.stat().st_size}))


if __name__ == "__main__":
    build()
